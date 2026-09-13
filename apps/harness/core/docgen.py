"""
Shared document builders — markdown-ish text → DOCX / PDF / XLSX bytes.

Extracted from the Second Brain export route so chat generation tools
(generate_document, generate_pdf, generate_spreadsheet) produce the same
quality of output: inline **bold** / *italic* / `code`, styled headings and
lists, and a consistent look across surfaces.

All builders are pure sync functions returning raw bytes; callers that need to
stay async should run them via loop.run_in_executor (see
api/routes/second_brain.py).
"""

import io
import re
from typing import Any, Dict, List


def build_docx(title: str, content: str, personal_note: str = "") -> bytes:
    """Markdown-ish text → DOCX bytes with inline bold/italic/code."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Render inline markdown (**bold**, *italic*, `code`) into runs on a paragraph.
    def add_inline(para, text):
        parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)', text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                para.add_run(part[2:-2]).bold = True
            elif part.startswith("*") and part.endswith("*"):
                para.add_run(part[1:-1]).italic = True
            elif part.startswith("`") and part.endswith("`"):
                run = para.add_run(part[1:-1])
                run.font.name = "Courier New"
                run.font.size = Pt(10)
            else:
                para.add_run(part)

    # Document title
    title_para = doc.add_heading(title, level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Personal note as italic block if present
    if personal_note:
        note_para = doc.add_paragraph()
        note_run = note_para.add_run(f"Note: {personal_note}")
        note_run.italic = True
        note_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        doc.add_paragraph()

    # Parse markdown line-by-line into docx
    lines = content.split("\n")
    for line in lines:
        # Headings
        if line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        # Unordered list items — inline-parsed
        elif re.match(r'^[-*+] ', line):
            add_inline(doc.add_paragraph(style="List Bullet"), line[2:])
        # Ordered list items — inline-parsed
        elif re.match(r'^\d+\. ', line):
            add_inline(doc.add_paragraph(style="List Number"), re.sub(r'^\d+\. ', '', line))
        # Horizontal rule
        elif line.strip() in ("---", "***", "___"):
            doc.add_paragraph("─" * 40)
        # Blank line
        elif line.strip() == "":
            pass
        else:
            # Normal paragraph — inline **bold**, *italic*, `code`
            add_inline(doc.add_paragraph(), line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def build_pdf(title: str, content: str, personal_note: str = "") -> bytes:
    """Markdown-ish text → styled PDF bytes (reportlab)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable

    buf = io.BytesIO()
    doc_pdf = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=25 * mm,
        rightMargin=25 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TARSTitle",
        parent=styles["Heading1"],
        fontSize=20,
        leading=26,
        spaceAfter=6,
        textColor=colors.HexColor("#1a1a1a"),
    )
    h1_style = ParagraphStyle("TARSH1", parent=styles["Heading1"], fontSize=16, leading=20, spaceAfter=4, spaceBefore=10)
    h2_style = ParagraphStyle("TARSH2", parent=styles["Heading2"], fontSize=13, leading=17, spaceAfter=3, spaceBefore=8)
    h3_style = ParagraphStyle("TARSH3", parent=styles["Heading3"], fontSize=11, leading=15, spaceAfter=2, spaceBefore=6)
    body_style = ParagraphStyle("TARSBody", parent=styles["Normal"], fontSize=10, leading=14, spaceAfter=4)
    note_style = ParagraphStyle("TARSNote", parent=styles["Normal"], fontSize=9, leading=13, textColor=colors.HexColor("#666666"), fontName="Helvetica-Oblique")
    bullet_style = ParagraphStyle("TARSBullet", parent=styles["Normal"], fontSize=10, leading=14, leftIndent=12, spaceAfter=2, bulletIndent=0)
    code_style = ParagraphStyle("TARSCode", parent=styles["Code"], fontSize=9, leading=12, backColor=colors.HexColor("#f4f4f4"))

    def inline_md(text):
        """Convert inline **bold** and *italic* to reportlab markup."""
        text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)
        text = re.sub(r'\*(.+?)\*', r'<i>\1</i>', text)
        text = re.sub(r'`(.+?)`', r'<font name="Courier">\1</font>', text)
        return text

    story = []
    story.append(Paragraph(title, title_style))
    story.append(Spacer(1, 4 * mm))

    if personal_note:
        story.append(Paragraph(f"Note: {personal_note}", note_style))
        story.append(Spacer(1, 3 * mm))

    lines = content.split("\n")
    for line in lines:
        if line.startswith("### "):
            story.append(Paragraph(inline_md(line[4:]), h3_style))
        elif line.startswith("## "):
            story.append(Paragraph(inline_md(line[3:]), h2_style))
        elif line.startswith("# "):
            story.append(Paragraph(inline_md(line[2:]), h1_style))
        elif re.match(r'^[-*+] ', line):
            story.append(Paragraph(f"• {inline_md(line[2:])}", bullet_style))
        elif re.match(r'^\d+\. ', line):
            num = re.match(r'^(\d+)\. ', line).group(1)
            line_text = re.sub(r'^\d+\. ', '', line)
            story.append(Paragraph(f"{num}. {inline_md(line_text)}", bullet_style))
        elif line.strip() in ("---", "***", "___"):
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#cccccc"), spaceAfter=4))
        elif line.strip() == "":
            story.append(Spacer(1, 3 * mm))
        else:
            story.append(Paragraph(inline_md(line), body_style))

    doc_pdf.build(story)
    buf.seek(0)
    return buf.read()


def build_xlsx(title: str, sheets: List[Dict[str, Any]]) -> bytes:
    """[{name, headers, rows}] → XLSX bytes (openpyxl).

    Light styling consistent with the other builders: bold header row, column
    widths sized to content, header row frozen. Cell values are stored as
    strings — the tool contract is string[][] so the model never has to guess
    at number coercion.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    # Remove the default sheet; every sheet comes from the caller
    wb.remove(wb.active)

    header_font = Font(bold=True)
    max_width = 60

    for i, sheet in enumerate(sheets or []):
        name = str(sheet.get("name") or f"Sheet {i + 1}")[:31] or f"Sheet {i + 1}"
        ws = wb.create_sheet(title=name)
        headers = [str(h) for h in (sheet.get("headers") or [])]
        rows = sheet.get("rows") or []

        widths = [len(h) for h in headers]
        if headers:
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font

        for row in rows:
            values = [str(v) if v is not None else "" for v in row]
            ws.append(values)
            for j, v in enumerate(values):
                if j < len(widths):
                    widths[j] = max(widths[j], len(v))
                else:
                    widths.append(len(v))

        for j, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(j + 1)].width = min(max_width, max(10, w + 2))

        if headers:
            ws.freeze_panes = "A2"

    # Never emit an empty workbook
    if not wb.sheetnames:
        ws = wb.create_sheet(title=str(title)[:31] or "Sheet 1")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
